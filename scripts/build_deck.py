"""Build docs/deck/Cadence.pptx from the report and the screenshots in docs/screenshots. Numbers on the slides are read from the database."""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

INK, INK2, CREAM, PAPER, LINE = RGBColor(0x1F, 0x20, 0x1D), RGBColor(0x5D, 0x5E, 0x58), RGBColor(0xF6, 0xF7, 0xEE), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE4, 0xE4, 0xE4)
LIVE, PAUSE, KILL = RGBColor(0x20, 0x74, 0x4C), RGBColor(0x9A, 0x5B, 0x00), RGBColor(0xC2, 0x3A, 0x2E)
W, H = Inches(13.333), Inches(7.5)
SHOTS = ROOT / "docs" / "screenshots"


def live_numbers() -> dict:
    from backend.analytics.rollups import stats
    from backend.core.db import tx

    with tx() as db:
        c = {cid: stats(db, cid) for cid in ("C1", "C2", "C3")}
        gold = db.q("select id, gold from prompt_versions where gold is not null and agent_key in ('Qualifier', 'Writer') and campaign_id = 'C1' order by id")
    return {"stats": c, "gold": {g["id"]: g["gold"][1] for g in gold}}


def slide(prs: Presentation, title: str, subtitle: str = "") -> object:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.35))
    band.fill.solid()
    band.fill.fore_color.rgb = CREAM
    band.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(30), True, INK
    if subtitle:
        st = s.shapes.add_textbox(Inches(0.6), Inches(0.85), Inches(12), Inches(0.4))
        sp = st.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size, sp.font.color.rgb = Pt(15), INK2
    return s


def bullets(s, items: list[str], left=0.6, top=1.7, width=12.0, size=18, gap=8) -> None:
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.3))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + it
        p.font.size, p.font.color.rgb = Pt(size), INK
        p.space_after = Pt(gap)


def box(s, x, y, w, h, text, fill=PAPER, color=INK, size=13, bold=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = 0.12
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    b.line.color.rgb = LINE
    tf = b.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size if i == 0 else size - 2), color, bold and i == 0
    return b


def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = INK2
    c.line.width = Pt(1.5)


def picture(s, name: str, left: float, top: float, width: float) -> None:
    path = SHOTS / name
    if path.exists():
        s.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def main() -> None:
    n = live_numbers()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    s = slide(prs, "Cadence", "An autonomous multi-channel SDR with a manager control plane")
    bullets(s, ["Launch a campaign, watch agents work it, pause it in one click.", "Live: buildathon-2026-production.up.railway.app", "Tech Contingent x DronaHQ Inter Guild Buildathon 2026",
                "Demo seller: Helix Agents (fictional). Every seeded record carries a DEMO chip."], top=2.2, size=22)

    s = slide(prs, "The problem", "How close is this to a real SDR working autonomously across channels?")
    bullets(s, ["Reps spend the day on research, copy and logging. Quality varies by rep and by day.", "Nobody can answer why a message went out.",
                "The brief: one integrated build. A control plane a manager operates, and agents that research, contact and follow up inside its campaigns.",
                "Judges weigh multi-channel intelligence, DronaHQ usage, personalisation without hallucination, and end-to-end capability."])

    s = slide(prs, "One SDR, not seven bots", "Five design rules, each with a visible artifact")
    bullets(s, ["One prospect record, one memory: a LinkedIn reply changes the next email.", "One planner owns channel choice.", "One policy gate before every send: ten ordered checks in plain code.",
                "One decision trace per action: prompt version, retrieved chunks, output, cost, gate result.", "One escalation path: the same object lands in the same inbox."])

    s = slide(prs, "Architecture", "Agents never send. Every send passes the Guardian, then a channel adapter.")
    box(s, 0.6, 2.0, 2.0, 0.9, "Manager\nDronaHQ Apps Studio")
    box(s, 3.2, 2.0, 2.0, 0.9, "Cadence API\nFastAPI + JWT")
    box(s, 5.8, 2.0, 2.2, 0.9, "Postgres\npgvector + full-text")
    box(s, 3.2, 3.6, 2.0, 0.9, "Worker\nSKIP LOCKED queue")
    box(s, 5.8, 3.6, 2.2, 0.9, "Guardian\ngate, conflicts, grounding", fill=CREAM)
    box(s, 8.6, 3.6, 2.2, 0.9, "Channels\nGmail, Twilio, LinkedIn (rep-assisted), voice")
    box(s, 3.2, 5.2, 2.0, 0.9, "Direct agents\nQualifier, Sequencer, Writer")
    box(s, 5.8, 5.2, 2.2, 0.9, "RAG\nhybrid retrieval, citations")
    box(s, 8.6, 5.2, 2.2, 0.9, "DronaHQ agents\nResearcher, Responder, Caller")
    for a in [(2.6, 2.45, 3.2, 2.45), (5.2, 2.45, 5.8, 2.45), (4.2, 2.9, 4.2, 3.6), (5.2, 4.05, 5.8, 4.05), (8.0, 4.05, 8.6, 4.05), (4.2, 4.5, 4.2, 5.2), (5.2, 5.65, 5.8, 5.65), (9.7, 5.2, 9.7, 4.5)]:
        arrow(s, *a)
    bullets(s, ["DronaHQ agents call our seven MCP tools. Callbacks carry a shared secret."], left=0.6, top=6.5, size=14)

    s = slide(prs, "The agents", "Six model-backed units and one deterministic supervisor")
    rows = [("Qualifier (ICP Fitment)", "Direct, fast model. Hard filters and the score are code"), ("Researcher", "DronaHQ agent with web search, enrichment and MCP tools"),
            ("Sequencer (Strategy and Follow-up)", "Direct, strong model. Code lists the legal channels first"), ("Writer (Personalisation)", "Direct, strong model. Grounding check rejects unsourced claims"),
            ("Responder (Conversation)", "DronaHQ agent. Opt-out and hard escalation rules run before any model"), ("Caller (Voice SDR)", "DronaHQ Voice with pre and post webhooks"),
            ("Guardian", "Plain code. No model can override it")]
    for i, (a, b) in enumerate(rows):
        box(s, 0.6, 1.6 + i * 0.75, 4.2, 0.6, a, fill=CREAM if i == 6 else PAPER, size=14, bold=True)
        tb = s.shapes.add_textbox(Inches(5.0), Inches(1.68 + i * 0.75), Inches(7.8), Inches(0.5))
        tb.text_frame.text = b
        tb.text_frame.paragraphs[0].font.size = Pt(15)

    s = slide(prs, "DronaHQ is core, and our code is real", "What runs where")
    bullets(s, ["Apps Studio: a native manager app with sign in, Command Center, campaigns, dashboard and kill switch, plus the full workspace embedded.", "The Researcher and Responder are DronaHQ Agentic platform agents. The Caller uses DronaHQ Voice.",
                "Our MCP server exposes seven tools to them: search_knowledge, get_timeline, save_research, propose_slots, book_meeting, create_escalation, set_classification.",
                "Our own code: state machine, gate, conflict engine, RAG, channels, evals, prompt versioning, auth.",
                "If a hosted agent is silent for 90 seconds, the step reruns on the direct provider and the feed says provider_fallback."])

    s = slide(prs, "The DronaHQ manager app", "Native Apps Studio screens on the live API (screens from a browser test against it)")
    picture(s, "dronahq_app_command_center.png", 0.5, 1.6, 6.1)
    picture(s, "dronahq_app_campaign_dashboard.png", 6.8, 1.6, 6.1)
    bullets(s, ["Sign in, Command Center with Stop all, campaign list with pause and resume, campaign dashboard with a switch per agent, and the full workspace embedded."], left=0.6, top=5.6, width=12.0, size=16)

    s = slide(prs, "The policy gate", "Ten ordered checks, one reason code each, idempotent")
    bullets(s, ["1 kill switch  2 campaign Live  3 agent on  4 channel on  5 not suppressed", "6 claim held by this campaign  7 frequency (one touch per 48 hours, four in 14 days)", "8 rep quota and active  9 campaign and channel daily caps  10 approval rule",
                "Per-prospect advisory lock, quota locks and a unique idempotency key: a race sends exactly one message.", "A Draft campaign refuses a real send attempt with campaign_not_live."], size=19)

    s = slide(prs, "Campaigns run side by side", "Pausing one writes one row and leaves the others running")
    picture(s, "flow_paused.png", 0.6, 1.6, 7.4)
    bullets(s, ["Three Live campaigns, one Draft.", "The worker's claim query joins each job to its own campaign.", "test_pause_isolation: pause C2, and within five seconds C1 and C3 keep writing activity while C2 writes none.", "Same guarantees for the agent and channel switches."], left=8.3, top=1.7, width=4.6, size=16)

    s = slide(prs, "Grounded, explainable personalisation", "Every claim traces to a source, and a decision trace shows why")
    picture(s, "prospects_sam-okafor.png", 0.6, 1.6, 7.4)
    bullets(s, ["Facts carry a source URL and a confidence.", "Retrieval is campaign-scoped and hybrid.", "Numbers must appear in their cited source. Banned phrases never pass.", "Two grounding failures fall back to a generic-safe draft or a human."], left=8.3, top=1.7, width=4.6, size=16)

    st = n["stats"]
    s = slide(prs, "Measurement", "Measured from the database, never typed in")
    bullets(s, [f"{cid}: {v['prospects']} prospects, {v['contacted']} contacted, reply rate {round(v['reply_rate'] * 100)}%, meetings {v['meetings']}, cost per qualified lead ${v['cost_per_qualified']:.2f}" for cid, v in st.items()]
            + ["Golden sets, 15 seeded cases per agent, scored by the runner: C1 Writer " + " and ".join(f"{k.split('-')[-1]} {v}" for k, v in n["gold"].items() if "Writer" in k) + "."], size=17)

    s = slide(prs, "Reliability", "Built to fail cleanly")
    bullets(s, ["99 tests against a real Postgres: every gate check, conflict cases K1 to K9, the send race, pause isolation, stop levels.", "LLM failure injection: invalid JSON, empty output, wrong schema, timeouts, 429s that wait as long as the provider asks. The Qualifier never qualifies on failure.",
                "Webhook secrets, JWT roles, injection, CORS, rate limits, and a log scan for emails.", "One error envelope, structured logs, typed exceptions, secrets in the environment, ALLOWED_RECIPIENTS on every send."], size=18)

    s = slide(prs, "What works, and what is honest", "Sandbox is an approved state and the badge tells the truth")
    bullets(s, ["Working live: campaigns, lifecycle, gate, conflicts, prompts and replay, RAG, evals, approvals, reps, analytics, MCP, and Gmail end to end (send, reply, classify, book).",
                "Partial: SMS (Twilio connected, first real text untested), voice (sandbox), and the hosted DronaHQ agents until they are switched on. Each falls back to sandbox or the direct provider.",
                "LinkedIn is rep-assisted by design: a person sends each note. Golden sets are small and the live-model scores depend on the provider's rate limit.", "Full list: docs/report/report.md"], size=18)

    s = slide(prs, "Try it yourself", "Judge card")
    bullets(s, ["1. Sign in with a demo chip. Create a campaign from the C1 template. Try Activate and read the checklist.", "2. Open Prompts on C3, edit the tone, save v3, and compare it with v2.", "3. Pause a campaign and watch Agent Activity.", "Live: buildathon-2026-production.up.railway.app", "Repo: github.com/kiranG18/Buildathon-2026"], size=20)

    out = ROOT / "docs" / "deck"
    out.mkdir(parents=True, exist_ok=True)
    prs.save(str(out / "Cadence.pptx"))
    print("saved", out / "Cadence.pptx", "with", len(prs.slides), "slides")


if __name__ == "__main__":
    main()
